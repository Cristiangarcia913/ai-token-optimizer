"""
Generates rich sample files across all formats to test and benchmark token reduction.
"""
import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
import json
import csv
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(SAMPLE_DIR, exist_ok=True)

def generate_html():
    path = os.path.join(SAMPLE_DIR, "webpage_article.html")
    content = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Informe de Mercado IA 2025</title>
    <style>
        body { font-family: 'Helvetica Neue', Arial, sans-serif; background-color: #f8fafc; color: #1e293b; margin: 0; padding: 20px; }
        .navbar { display: flex; justify-content: space-between; background: #0f172a; color: white; padding: 15px; }
        .hero { padding: 40px; background: linear-gradient(135deg, #3b82f6, #1d4ed8); color: white; text-align: center; }
        .card { background: white; border-radius: 8px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); padding: 20px; margin-bottom: 20px; }
        .footer { background: #0f172a; color: #94a3b8; text-align: center; padding: 30px; font-size: 12px; }
    </style>
    <script>
        console.log("Analytics tracker initialized: " + Date.now());
        window.addEventListener('load', function() {
            fetch('/api/telemetry', { method: 'POST', body: JSON.stringify({ page: 'market_report', referrer: document.referrer }) });
        });
    </script>
</head>
<body>
    <div class="cookie-banner" id="cookie-consent">
        <p>Utilizamos cookies y tecnologías similares para mejorar su experiencia de usuario.</p>
        <button onclick="acceptCookies()">Aceptar todas</button>
    </div>

    <nav class="navbar">
        <div class="logo">TechTrends Global</div>
        <ul class="nav-links">
            <li><a href="#home">Inicio</a></li>
            <li><a href="#reports">Informes</a></li>
            <li><a href="#contact">Contacto</a></li>
        </ul>
    </nav>

    <main class="container">
        <article class="card">
            <h1>El Estado de la Inteligencia Artificial Generativa en 2025</h1>
            <p>La adopción de modelos de lenguaje grande (LLMs) ha alcanzado un punto de inflexión sin precedentes en el entorno empresarial. Más del 78% de las organizaciones del Fortune 500 han integrado asistentes inteligentes en sus flujos operativos de soporte, desarrollo de software y análisis de datos.</p>
            
            <h2>Principales Factores de Crecimiento</h2>
            <ul>
                <li><strong>Eficiencia Operativa:</strong> Automatización del 60% de tareas repetitivas de redacción y análisis documental.</li>
                <li><strong>Reducción de Costes:</strong> Disminución del 45% en tiempos de respuesta al cliente.</li>
                <li><strong>Capacidades Multimodales:</strong> Análisis directo de imágenes, diagramas de arquitectura y transcripciones de audio.</li>
            </ul>

            <h2>Comparativa de Despliegue Empresarial</h2>
            <table>
                <thead>
                    <tr>
                        <th>Sector Industrial</th>
                        <th>Nivel de Adopción (%)</th>
                        <th>ROI Promedio Anual</th>
                        <th>Caso de Uso Principal</th>
                    </tr>
                </thead>
                <tbody>
                    <tr>
                        <td>Finanzas & Banca</td>
                        <td>84%</td>
                        <td>3.4x</td>
                        <td>Detección de fraude y extracción documental</td>
                    </tr>
                    <tr>
                        <td>Salud & Farmacia</td>
                        <td>71%</td>
                        <td>2.8x</td>
                        <td>Resumen de historiales clínicos y triaje</td>
                    </tr>
                    <tr>
                        <td>Comercio Electrónico</td>
                        <td>92%</td>
                        <td>4.1x</td>
                        <td>Atención al cliente 24/7 y recomendaciones</td>
                    </tr>
                </tbody>
            </table>

            <blockquote>
                "El valor diferencial ya no radica en disponer de un modelo fundacional, sino en la calidad y síntesis de la información que se le suministra en contexto."
            </blockquote>
        </article>
    </main>

    <footer class="footer">
        <p>© 2025 TechTrends Media Inc. Todos los derechos reservados. Términos de Servicio | Política de Privacidad.</p>
    </footer>
</body>
</html>
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path

def generate_json():
    path = os.path.join(SAMPLE_DIR, "enterprise_customers.json")
    customers = []
    tiers = ["Enterprise", "Professional", "Standard"]
    regions = ["EMEA", "LATAM", "APAC", "North America"]
    statuses = ["active", "trialing", "renewed"]

    for i in range(1, 41):
        customers.append({
            "account_id": f"ACC-2025-{i:04d}",
            "company_name": f"Global Solutions Tech {i} S.L.",
            "contact_email": f"director.operaciones{i}@globalsolutions{i}.com",
            "tier": tiers[i % len(tiers)],
            "region": regions[i % len(regions)],
            "monthly_recurring_revenue": 1500 + (i * 250),
            "contract_status": statuses[i % len(statuses)],
            "license_seats": 20 + (i * 5),
            "cloud_provider": "Google Cloud",
            "support_priority": "High" if i % 2 == 0 else "Normal"
        })

    with open(path, "w", encoding="utf-8") as f:
        json.dump(customers, f, indent=4, ensure_ascii=False)
    return path

def generate_csv():
    path = os.path.join(SAMPLE_DIR, "sales_pipeline.csv")
    headers = ["Deal ID", "Account Name", "Stage", "Deal Value ($)", "Probability (%)", "Close Date", "Owner", "", ""]
    rows = [headers]
    stages = ["Qualification", "Proposal Sent", "Contract Review", "Closed Won"]
    owners = ["Carlos Ruiz", "Elena Martin", "David Garcia", "Sofia Lopez"]

    for i in range(1, 35):
        rows.append([
            f"DEAL-{1000 + i}",
            f"Corporacion Alfa {i}",
            stages[i % len(stages)],
            str(12000 + (i * 3500)),
            f"{25 * ((i % 4) + 1)}%",
            f"2025-11-{(i % 28) + 1:02d}",
            owners[i % len(owners)],
            "", # empty trailing columns common in Excel exports
            ""
        ])

    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)
    return path

def generate_code():
    path = os.path.join(SAMPLE_DIR, "service_auth.py")
    content = '''# Copyright (c) 2024-2025 Acme Software Corporation
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#
# System Identification: AUTH-SERVICE-SEC-092
# Module: auth_token_provider.py
# Author: Architecture Security Guild

import hmac
import hashlib
import time
from typing import Optional, Dict, Any



def verify_jwt_token(token: str, secret_key: str) -> bool:
    if not token or not secret_key:
        return False
        
    parts = token.split(".")
    if len(parts) != 3:
        return False
        
    header, payload, signature = parts
    expected = hmac.new(
        secret_key.encode("utf-8"),
        f"{header}.{payload}".encode("utf-8"),
        hashlib.sha256
    ).hexdigest()
    
    return hmac.compare_digest(expected, signature)



class SecuritySessionManager:
    def __init__(self, ttl_seconds: int = 3600):
        self.ttl = ttl_seconds
        self.sessions: Dict[str, float] = {}

    def register_session(self, user_id: str) -> None:
        self.sessions[user_id] = time.time()

    def is_valid(self, user_id: str) -> bool:
        created = self.sessions.get(user_id)
        if not created:
            return False
        return (time.time() - created) < self.ttl
'''
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path

def generate_docx():
    path = os.path.join(SAMPLE_DIR, "strategic_plan.docx")
    doc = docx.Document()
    doc.add_heading("Plan Estratégico de Transformación Digital 2025", 0)
    
    p = doc.add_paragraph("Este documento describe los objetivos prioritarios para la adopción de inteligencia artificial generativa en toda la infraestructura de la empresa.")
    
    doc.add_heading("1. Objetivos del Trimestre Q1", level=1)
    doc.add_paragraph("Consolidación de la base de conocimiento y reducción de latencia en modelos conversacionales.", style='List Bullet')
    doc.add_paragraph("Implementación de un optimizador de tokens para reducir los costes de inferencia en un 50%.", style='List Bullet')
    doc.add_paragraph("Auditoría de cumplimiento normativo y privacidad de datos de clientes.", style='List Bullet')

    doc.add_heading("2. Asignación Presupuestaria", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Iniciativa"
    hdr_cells[1].text = "Presupuesto (€)"
    hdr_cells[2].text = "Impacto Estimado"

    data = [
        ("Optimización de Tokens y RAG", "45.000 €", "Ahorro recurrente del 65% en API"),
        ("Despliegue Local de LLMs (vLLM)", "80.000 €", "Soberanía de datos y baja latencia"),
        ("Formación del Personal Técnico", "25.000 €", "Mejora de productividad interna")
    ]
    for init, budget, impact in data:
        row_cells = table.add_row().cells
        row_cells[0].text = init
        row_cells[1].text = budget
        row_cells[2].text = impact

    doc.save(path)
    return path

def generate_pdf():
    path = os.path.join(SAMPLE_DIR, "audit_report.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    
    # Page 1
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 750, "CONFIDENCIAL - AUDITORÍA INTERNA CORPORATIVA 2025")
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, 680, "Informe de Auditoría de Sistemas y Seguridad")
    c.setFont("Helvetica", 11)
    c.drawString(50, 640, "La presente auditoría evalúa el cumplimiento de las normativas de protección")
    c.drawString(50, 620, "de datos y la eficiencia en el procesamiento de modelos de inteli-")
    c.drawString(50, 600, "gencia artificial en el entorno cloud de producción.")
    c.drawString(50, 560, "Hallazgos Principales:")
    c.drawString(60, 530, "- Se detectó un sobrecoste de tokens por inclusión de cabeceras redundantes.")
    c.drawString(60, 500, "- La latencia de respuesta se redujo un 40% al limpiar documentos Markdown.")
    c.setFont("Helvetica", 9)
    c.drawString(250, 50, "Página 1 de 2")
    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 750, "CONFIDENCIAL - AUDITORÍA INTERNA CORPORATIVA 2025")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 680, "Recomendaciones de Optimización")
    c.setFont("Helvetica", 11)
    c.drawString(50, 640, "Se recomienda estandarizar el preprocesamiento de todos los archivos antes")
    c.drawString(50, 620, "de ser inyectados en la ventana de contexto de los agentes inteligentes.")
    c.setFont("Helvetica", 9)
    c.drawString(250, 50, "Página 2 de 2")
    c.showPage()
    
    c.save()
    return path

def generate_excel():
    import openpyxl
    path = os.path.join(SAMPLE_DIR, "financial_projections.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1_Q4_Projections"

    headers = ["Quarter", "Revenue ($k)", "OpEx ($k)", "Net Margin (%)", "New ARR ($k)", "Headcount", "", ""]
    ws.append(headers)

    quarters = ["Q1-2025", "Q2-2025", "Q3-2025", "Q4-2025"]
    for i, q in enumerate(quarters, 1):
        ws.append([q, 2400 + (i * 350), 1800 + (i * 120), f"{18 + (i * 2)}%", 450 + (i * 90), 85 + (i * 10), "", ""])

    # Add trailing empty rows (typical in dirty spreadsheet exports)
    for _ in range(5):
        ws.append(["", "", "", "", "", "", "", ""])

    wb.save(path)
    return path

if __name__ == "__main__":
    generate_html()
    generate_json()
    generate_csv()
    generate_excel()
    generate_code()
    generate_docx()
    generate_pdf()
    print("✅ All sample files generated in:", SAMPLE_DIR)
